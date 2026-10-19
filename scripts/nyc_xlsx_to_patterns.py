"""Convert NYC BOE 2021 primary CVR xlsx files (PE2021_CVR_Final.zip) into a pattern TSV for one RCV contest.
Usage: python3 scripts/nyc_xlsx_to_patterns.py <cvr_dir> "<column prefix>" <out.tsv>
Cells are candidacy IDs, 'undervote', 'overvote' or 'Write-in'. IDs are mapped to names with 2021P_CandidacyID_To_Name.xlsx.
Ballots with no mark in any rank of the contest (voters of other parties / no vote in this race) are dropped and counted."""
import openpyxl, glob, os, sys, collections, json, warnings
warnings.filterwarnings('ignore')
cvr_dir, prefix, out = sys.argv[1:4]
names = {}
ws = openpyxl.load_workbook(f'{cvr_dir}/2021P_CandidacyID_To_Name.xlsx', read_only=True).active; ws.reset_dimensions()
for i, row in enumerate(ws.iter_rows(values_only=True)):
    if i and row[0] is not None: names[str(row[0]).strip()] = str(row[1]).strip()
counts = collections.Counter(); stats = {}; dropped = 0
for f in sorted(glob.glob(f'{cvr_dir}/2021P*V*.xlsx')):
    if 'CandidacyID' in f: continue
    ws = openpyxl.load_workbook(f, read_only=True).active; ws.reset_dimensions()
    it = ws.iter_rows(values_only=True); header = next(it)
    cols = [i for i, c in enumerate(header) if c and str(c).startswith(prefix)]
    if not cols: stats[os.path.basename(f)] = 'no contest columns'; continue
    cols.sort(key=lambda i: str(header[i]))  # 'Choice 1 of 5' ... 'Choice 5 of 5'
    n = 0
    for row in it:
        if row is None or row[0] is None: continue
        marks = []
        for i in cols:
            v = '' if row[i] is None else str(row[i]).strip()
            if v in ('', 'undervote'): marks.append('undervote')
            elif v == 'overvote': marks.append('overvote')
            elif v == 'Write-in': marks.append('Write-in')
            else:
                if v not in names: raise SystemExit(f'unknown candidacy id {v} in {f}')
                marks.append(names[v])
        if all(m == 'undervote' for m in marks): dropped += 1; continue
        counts[tuple(marks)] += 1; n += 1
    stats[os.path.basename(f)] = n
with open(out, 'w') as o:
    for k, c in sorted(counts.items(), key=lambda kv: (-kv[1], kv[0])): o.write(str(c) + '\t' + '\t'.join(k) + '\n')
print(json.dumps({'files': stats, 'ballots': sum(counts.values()), 'droppedNoMark': dropped, 'patterns': len(counts)}))
